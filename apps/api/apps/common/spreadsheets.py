"""
Reading and writing the file a shop already keeps its catalogue in.

A merchant moving to this platform has their products in a spreadsheet, and
they will want their reports back out in one. Both directions go through here
so that a column added to an export automatically becomes a column the importer
understands, and the two cannot drift apart.

CSV and XLSX only. They are what the shops actually have — Excel, Google
Sheets, and whatever their old till exported — and both are readable without
any system library.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from decimal import Decimal
from typing import Any

CSV = "csv"
XLSX = "xlsx"

FORMATS = (CSV, XLSX)

CONTENT_TYPES = {
    CSV: "text/csv; charset=utf-8",
    XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass(frozen=True)
class Column:
    """One column of an import/export sheet.

    `key` is the field the application uses; `header` is what the merchant
    reads. They are separate because the header is translated copy and the key
    is a contract — renaming the heading must not silently break every saved
    spreadsheet a shop imports from.
    """

    key: str
    header: str
    #: Shown in the downloadable template so the format is obvious from an
    #: example rather than from documentation nobody opens.
    example: str = ""
    help_text: str = ""


def render(*, columns: list[Column], rows: list[dict[str, Any]], fmt: str) -> bytes:
    """Serialise rows into a spreadsheet."""
    if fmt == XLSX:
        return _render_xlsx(columns=columns, rows=rows)
    return _render_csv(columns=columns, rows=rows)


def _cell(value: Any) -> Any:
    """Flatten a value into something a spreadsheet cell can hold."""
    if value is None:
        return ""
    if isinstance(value, bool):
        # Not "True"/"False": the merchant edits this by hand and sends it back.
        return "sim" if value else "nao"
    if isinstance(value, Decimal):
        return str(value)
    return value


def _render_csv(*, columns: list[Column], rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow([column.header for column in columns])
    for row in rows:
        writer.writerow([_cell(row.get(column.key)) for column in columns])

    # A BOM, because Excel in a Portuguese locale opens a UTF-8 CSV as Latin-1
    # otherwise and turns every "ã" into mojibake. The shop would conclude the
    # export is broken, and they would be right to.
    return buffer.getvalue().encode("utf-8-sig")


def _render_xlsx(*, columns: list[Column], rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="8C1425")

    sheet.append([column.header for column in columns])
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        # Roughly the header's own width, so nothing is hidden behind "####"
        # when the file is opened.
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(column.header) + 4)

    for row in rows:
        sheet.append([_cell(row.get(column.key)) for column in columns])

    # The header stays put while a merchant scrolls a thousand products.
    sheet.freeze_panes = "A2"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse(*, content: bytes, filename: str, columns: list[Column]) -> list[dict[str, Any]]:
    """Read a spreadsheet back into rows keyed by column `key`.

    Headers are matched leniently — case and surrounding space are ignored, and
    a column's own `key` is accepted as well as its header. A shop that exported
    from here, edited in Excel and saved as CSV should not be told their file is
    unrecognisable because a heading gained a capital letter.
    """
    table = _read_xlsx(content) if filename.lower().endswith(".xlsx") else _read_csv(content)
    if not table:
        return []

    lookup = {}
    for column in columns:
        lookup[_normalise(column.header)] = column.key
        lookup[_normalise(column.key)] = column.key

    headers = [lookup.get(_normalise(str(cell)), "") for cell in table[0]]

    rows = []
    for values in table[1:]:
        row: dict[str, Any] = {}
        for header, value in zip(headers, values, strict=False):
            if header:
                row[header] = value.strip() if isinstance(value, str) else value
        # A trailing blank line is what a spreadsheet leaves behind, not a row
        # the merchant meant to send.
        if any(str(value or "").strip() for value in row.values()):
            rows.append(row)

    return rows


def _normalise(value: str) -> str:
    return " ".join(value.strip().lower().split())


def _read_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")

    # Excel writes `;` in a Portuguese locale and `,` elsewhere; both turn up.
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","

    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active

    return [
        ["" if cell is None else cell for cell in row] for row in sheet.iter_rows(values_only=True)
    ]


def download(*, columns: list[Column], rows: list[dict[str, Any]], stem: str, fmt: str) -> Any:
    """Render rows and hand them over as a file download.

    The filename carries the date, because a merchant ends up with several of
    these in their downloads folder and `produtos (3).xlsx` tells them nothing.
    """
    from django.http import HttpResponse
    from django.utils import timezone
    from django.utils.text import slugify

    # The query parameter is `fmt`, not `format`: DRF reserves `format` for
    # content negotiation, so `?format=csv` is answered with a 404 for a
    # renderer that does not exist rather than reaching the view at all.
    chosen = fmt if fmt in FORMATS else XLSX
    content = render(columns=columns, rows=rows, fmt=chosen)
    filename = f"{slugify(stem)}-{timezone.localdate():%Y-%m-%d}.{chosen}"

    response = HttpResponse(content, content_type=CONTENT_TYPES[chosen])
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    # A browser cannot read the filename off a cross-origin download without
    # this, and the storefront and the API sit on different origins.
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    return response
